"""
Excel sheets for the escalation workflow and the supply side:
  Escalation Queue   prioritised, owned, dated actions with a status dropdown
  Inventory          reorder-point / safety-stock policy table
  Forecast           actual vs forecast network demand + category accuracy
  Supplier Scorecard OTIF by supplier with chart
"""

from __future__ import annotations

import pandas as pd
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

import config
from modules.excel_reporter import (
    CLR, GRN_FILL, INT_FMT, PCT_FMT, RED_FILL, YEL_FILL, _auto_width, _cell, _hdr,
)

SHEET_QUEUE = "Escalation Queue"
SHEET_INVENTORY = "Inventory"
SHEET_FORECAST = "Forecast"
SHEET_SUPPLIER = "Supplier Scorecard"
USD_FMT = '"$"#,##0'

WORKFLOW_STATES = "OPEN,IN PROGRESS,BLOCKED,DONE"


def _title(ws, text: str, span: str) -> None:
    ws.merge_cells(span)
    ws["A1"].value = text
    ws["A1"].font = Font(name="Calibri", bold=True, size=12, color=CLR["header_bg"])
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 24


def _status_cf(ws, rng: str, red: list[str], yellow: list[str], green: list[str]) -> None:
    for vals, fill in ((red, RED_FILL), (yellow, YEL_FILL), (green, GRN_FILL)):
        for v in vals:
            ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=[f'"{v}"'], fill=fill))


def build_queue_sheet(ws, queue: pd.DataFrame, report_date: str) -> None:
    ws.sheet_view.showGridLines = False
    _title(ws, f"Escalation Queue - report date {report_date}", "A1:N1")
    heads = ["ID", "Priority", "Score", "Domain", "Entity", "Issue", "Metric", "Value",
             "Benchmark", "Est. $ Impact", "Owner", "Due Date", "Recommended Action", "Workflow Status"]
    for i, h in enumerate(heads, 1):
        _hdr(ws, 2, i, h)
    for r, row in enumerate(queue.itertuples(index=False), 3):
        vals = [row.id, row.priority, row.severity_score, row.domain, row.entity, row.issue,
                row.metric, row.value, row.benchmark, row.est_cost_usd, row.owner, row.due_date,
                row.action, row.status]
        fmts = [None, None, "0", None, None, None, None, PCT_FMT, PCT_FMT, USD_FMT, None, None, None, None]
        for c, (v, f) in enumerate(zip(vals, fmts), 1):
            _cell(ws, r, c, v, fmt=f, bold=(c in (1, 2)), align="left" if c in (4, 5, 6, 11, 13) else "center")
            if c in (6, 13):
                ws.cell(row=r, column=c).alignment = Alignment(wrap_text=True, vertical="top")
    last = 2 + len(queue)
    _status_cf(ws, f"B3:B{max(last, 3)}", ["P1"], ["P2"], ["P3"])
    _status_cf(ws, f"N3:N{max(last, 3)}", ["BLOCKED"], ["IN PROGRESS"], ["DONE"])
    dv = DataValidation(type="list", formula1=f'"{WORKFLOW_STATES}"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add(f"N3:N{max(last, 3)}")
    ws.auto_filter.ref = f"A2:N{last}"
    ws.freeze_panes = "C3"
    for i, w in enumerate([9, 9, 7, 12, 30, 62, 16, 10, 11, 14, 26, 12, 62, 16], 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def build_inventory_sheet(ws, policy: pd.DataFrame) -> None:
    ws.sheet_view.showGridLines = False
    _title(ws, "Inventory Position & Reorder Policy (SS = z*sqrt(LT*sd_d^2 + d^2*sd_LT^2); ROP = d*LT + SS)", "A1:Q1")
    cols = [("node", "Node", None), ("category", "Category", None), ("supplier_id", "Supplier", None),
            ("on_hand", "On Hand", INT_FMT), ("on_order", "On Order", INT_FMT),
            ("avg_daily_demand", "Avg Daily Demand", INT_FMT), ("days_of_supply", "Days of Supply", "0.0"),
            ("fill_rate", "Fill Rate (90d)", PCT_FMT), ("lead_time_mean", "Lead Time (d)", "0.0"),
            ("lead_time_sd", "Lead Time SD", "0.0"), ("safety_stock", "Safety Stock", INT_FMT),
            ("reorder_point", "Reorder Point", INT_FMT), ("current_policy_rop", "Current Policy ROP", INT_FMT),
            ("rop_gap_units", "ROP Gap (units)", INT_FMT), ("recommended_order_qty", "Rec. Order Qty", INT_FMT),
            ("lost_margin_usd", "Lost Margin $ (90d)", USD_FMT), ("status", "Status", None)]
    for i, (_, h, _) in enumerate(cols, 1):
        _hdr(ws, 2, i, h)
    for r, (_, row) in enumerate(policy.iterrows(), 3):
        for c, (k, _, f) in enumerate(cols, 1):
            v = row[k]
            v = None if pd.isna(v) else (v.item() if hasattr(v, "item") else v)
            _cell(ws, r, c, v, fmt=f, align="left" if c <= 3 else "center")
    last = 2 + len(policy)
    _status_cf(ws, f"Q3:Q{last}", ["STOCKED OUT"], ["REORDER", "EXCESS"], ["OK"])
    ws.conditional_formatting.add(f"H3:H{last}", CellIsRule(operator="lessThan", formula=["0.9"], fill=RED_FILL))
    ws.conditional_formatting.add(f"H3:H{last}", CellIsRule(operator="lessThan", formula=["0.97"], fill=YEL_FILL))
    ws.auto_filter.ref = f"A2:Q{last}"
    ws.freeze_panes = "D3"
    _auto_width(ws, min_w=10, max_w=24)


def build_forecast_sheet(ws, daily: pd.DataFrame, forecast: pd.DataFrame, accuracy: pd.DataFrame) -> None:
    ws.sheet_view.showGridLines = False
    _title(ws, "Demand Forecast - network total, actual vs forecast", "A1:H1")

    # category accuracy table
    for i, h in enumerate(["Category", "Median WAPE (backtest)", "Median gain vs naive", "Best method",
                           f"Forecast units ({config.FORECAST_HORIZON_DAYS}d)"], 1):
        _hdr(ws, 3, i, h)
    cat = (accuracy.groupby("category")
           .agg(wape=("best_wape", "median"), gain=("improvement_vs_naive", "median"),
                method=("best_method", lambda s: s.mode().iat[0]),
                units=("forecast_total", "sum")).reset_index())
    for r, row in enumerate(cat.itertuples(index=False), 4):
        _cell(ws, r, 1, row.category, bold=True, align="left")
        _cell(ws, r, 2, float(row.wape), fmt=PCT_FMT)
        _cell(ws, r, 3, float(row.gain), fmt=PCT_FMT)
        _cell(ws, r, 4, row.method)
        _cell(ws, r, 5, float(row.units), fmt=INT_FMT)
    r0 = 4 + len(cat) + 2
    ws.cell(row=r0 - 1, column=1,
            value="WAPE = sum|actual - forecast| / sum(actual) on the last 14 days held out. Method chosen per series by backtest.") \
        .font = Font(name="Calibri", size=8, color=CLR["text_muted"])

    # network daily table
    hist = (daily.assign(date=pd.to_datetime(daily["date"]))
            .groupby("date")["demand_units"].sum().tail(28))
    fc = (forecast.assign(date=pd.to_datetime(forecast["date"]))
          .groupby("date")["forecast_units"].sum())
    for i, h in enumerate(["Date", "Actual units", "Forecast units"], 1):
        _hdr(ws, r0, i, h)
    r = r0
    for d, v in hist.items():
        r += 1
        _cell(ws, r, 1, d.to_pydatetime(), fmt="yyyy-mm-dd", align="left")
        _cell(ws, r, 2, float(v), fmt=INT_FMT)
        _cell(ws, r, 3, None)
    for d, v in fc.items():
        r += 1
        _cell(ws, r, 1, d.to_pydatetime(), fmt="yyyy-mm-dd", align="left")
        _cell(ws, r, 2, None)
        _cell(ws, r, 3, float(v), fmt=INT_FMT)

    ch = LineChart()
    ch.title = "Network demand: last 28 days actual, next 14 days forecast"
    ch.height, ch.width = 8, 18
    ch.add_data(Reference(ws, min_col=2, max_col=3, min_row=r0, max_row=r), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=r0 + 1, max_row=r))
    ch.x_axis.number_format = "mm-dd"
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ws.add_chart(ch, "G3")
    for i, w in enumerate([24, 22, 20, 14, 20], 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def build_supplier_sheet(ws, sc: pd.DataFrame) -> None:
    ws.sheet_view.showGridLines = False
    _title(ws, f"Supplier Scorecard - OTIF target {config.OTIF_TARGET:.0%} (on-time AND in-full, per PO)", "A1:O1")
    cols = [("supplier_id", "Supplier", None), ("category", "Category", None),
            ("pos_received", "POs Received", INT_FMT), ("on_time_rate", "On-Time", PCT_FMT),
            ("in_full_rate", "In-Full", PCT_FMT), ("otif", "OTIF", PCT_FMT),
            ("lead_time_mean", "Lead Time (d)", "0.0"), ("lead_time_sd", "Lead Time SD", "0.0"),
            ("avg_days_late", "Avg Days Late", "0.0"), ("fill_rate", "Fill Rate", PCT_FMT),
            ("spend", "Spend $", USD_FMT), ("shortfall_value", "Shortfall $", USD_FMT),
            ("open_pos", "Open POs", INT_FMT), ("otif_gap_pts", "OTIF Gap (pts)", "0.0"), ("status", "Status", None)]
    for i, (_, h, _) in enumerate(cols, 1):
        _hdr(ws, 2, i, h)
    for r, (_, row) in enumerate(sc.iterrows(), 3):
        for c, (k, _, f) in enumerate(cols, 1):
            v = row[k]
            v = v.item() if hasattr(v, "item") else v
            _cell(ws, r, c, v, fmt=f, bold=(c == 1), align="left" if c <= 2 else "center")
    last = 2 + len(sc)
    _status_cf(ws, f"O3:O{last}", ["ESCALATE"], ["WATCH"], ["OK"])
    ws.conditional_formatting.add(f"F3:F{last}", CellIsRule(operator="lessThan", formula=[str(config.OTIF_WATCH)], fill=RED_FILL))
    ws.conditional_formatting.add(f"F3:F{last}", CellIsRule(operator="lessThan", formula=[str(config.OTIF_TARGET)], fill=YEL_FILL))
    ws.auto_filter.ref = f"A2:O{last}"
    ws.freeze_panes = "C3"
    _auto_width(ws, min_w=10, max_w=22)

    ch = BarChart()
    ch.type = "bar"
    ch.title = "OTIF by supplier"
    ch.height, ch.width = 8, 16
    ch.add_data(Reference(ws, min_col=6, min_row=2, max_row=last), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=3, max_row=last))
    ch.y_axis.number_format = "0%"
    ch.y_axis.scaling.min, ch.y_axis.scaling.max = 0, 1
    ch.legend = None
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ws.add_chart(ch, f"A{last + 3}")
