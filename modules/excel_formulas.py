"""
Excel sheets that stay LIVE inside Excel (formulas, not pasted values):

  Data_Weekly    weekly facts by node x carrier x category (the pivot's source)
  Formula Pivot  dropdown-driven SUMIFS / AVERAGEIFS / RANK / INDEX-MATCH tables,
                 conditional formatting driven by editable threshold cells,
                 a native line chart, and a reconciliation check against Python
  Assumptions    every threshold / weight / cost used, and what is simulated

An analyst can change the Node / Carrier / Category dropdowns or the two
threshold cells and the whole sheet recalculates with no Python and no macro.
"""

from __future__ import annotations

import pandas as pd
from openpyxl.chart import LineChart, Reference
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

import config
from modules.excel_reporter import (
    CLR, GRN_FILL, INT_FMT, PCT_FMT, RED_FILL, SUB_FILL, YEL_FILL,
    _auto_width, _cell, _hdr,
)

SHEET_PIVOT = "Formula Pivot"
SHEET_DATA = "Data_Weekly"
SHEET_ASSUMP = "Assumptions"

DEFECT_COLS = list(config.DEFECT_WEIGHTS)
DEFECT_LABEL = {
    "failed_pickup": "Failed Pickup", "late_delivery": "Late Delivery",
    "missing_package": "Missing Package", "damaged_goods": "Damaged Goods",
    "wrong_item": "Wrong Item", "inventory_discrepancy": "Inv Discrepancy",
}
# Data_Weekly column letters
DW = {"week": "A", "node": "B", "carrier": "C", "category": "D",
      "ship": "E", "defects": "F", "days": "M"}
PIVOT_DEFECT_ORDER = ["late_delivery", "missing_package", "damaged_goods",
                      "wrong_item", "inventory_discrepancy", "failed_pickup"]
for _i, _c in enumerate(PIVOT_DEFECT_ORDER):
    DW[_c] = get_column_letter(7 + _i)       # G..L


def weekly_data(df: pd.DataFrame) -> pd.DataFrame:
    d = df.copy()
    dt = pd.to_datetime(d["date"])
    d["week_start"] = (dt - pd.to_timedelta(dt.dt.weekday, unit="D")).dt.normalize()
    days = d.groupby("week_start")["date"].nunique().rename("days_in_week")
    g = (d.groupby(["week_start", "node", "carrier", "category"], as_index=False)
          [["total_shipments", "total_defects", *PIVOT_DEFECT_ORDER]].sum())
    return g.merge(days, on="week_start").sort_values(
        ["week_start", "node", "carrier", "category"]).reset_index(drop=True)


def build_data_sheet(ws, wk: pd.DataFrame) -> int:
    """Write Data_Weekly; return the last data row."""
    heads = ["week_start", "node", "carrier", "category", "total_shipments", "total_defects",
             *PIVOT_DEFECT_ORDER, "days_in_week"]
    for i, h in enumerate(heads, 1):
        _hdr(ws, 1, i, h)
    cols = ["week_start", "node", "carrier", "category", "total_shipments", "total_defects",
            *PIVOT_DEFECT_ORDER, "days_in_week"]
    for r, row in enumerate(wk[cols].itertuples(index=False), 2):
        for c, v in enumerate(row, 1):
            if c == 1:
                v = v.to_pydatetime()
            elif hasattr(v, "item"):
                v = v.item()
            cell = ws.cell(row=r, column=c, value=v)
            if c == 1:
                cell.number_format = "yyyy-mm-dd"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(heads))}{1 + len(wk)}"
    for i in range(1, len(heads) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 16
    return 1 + len(wk)


def _rng(col: str, last: int) -> str:
    return f"{SHEET_DATA}!${col}$2:${col}${last}"


def _sumifs(sum_col: str, last: int, crits: list[tuple[str, str]]) -> str:
    parts = ",".join(f"{_rng(c, last)},{ref}" for c, ref in crits)
    return f"=SUMIFS({_rng(sum_col, last)},{parts})"


def build_pivot_sheet(ws, wk: pd.DataFrame, data_last: int, python_total_defects: int) -> None:
    ws.sheet_view.showGridLines = False
    nodes = sorted(wk["node"].unique())
    carriers = sorted(wk["carrier"].unique())
    cats = sorted(wk["category"].unique())
    weeks = sorted(wk["week_start"].unique())

    ws.merge_cells("A1:M1")
    ws["A1"].value = "Formula Pivot - live Excel formulas (change the yellow cells)"
    ws["A1"].font = Font(name="Calibri", bold=True, size=13, color=CLR["header_bg"])
    ws["A2"].value = ("Every number below is a SUMIFS / AVERAGEIFS / RANK / INDEX-MATCH formula over the "
                      "Data_Weekly sheet. No macro, no Python needed to re-slice.")
    ws["A2"].font = Font(name="Calibri", size=9, color=CLR["text_muted"])

    # selectors
    inputs = [("Node", "B4", nodes), ("Carrier", "B5", carriers), ("Category", "B6", cats)]
    for i, (label, ref, opts) in enumerate(inputs):
        row = 4 + i
        ws.cell(row=row, column=1, value=label).font = Font(name="Calibri", bold=True, size=10)
        c = ws[ref]
        c.value = "All"
        c.fill = YEL_FILL
        c.font = Font(name="Calibri", bold=True, size=10, color="1D4ED8")
        dv = DataValidation(type="list", formula1='"' + ",".join(["All", *opts]) + '"',
                            allow_blank=False)
        ws.add_data_validation(dv)
        dv.add(c)
        # criteria helper: "All" -> wildcard
        h = ws.cell(row=row, column=3, value=f'=IF(B{row}="All","*",B{row})')
        h.font = Font(name="Courier New", size=9, color=CLR["text_muted"])
    ws["C3"].value = "criteria used"
    ws["C3"].font = Font(name="Calibri", size=8, color=CLR["text_muted"])

    ws["E4"].value, ws["E5"].value = "Escalate rate >=", "Watch rate >="
    for ref in ("E4", "E5"):
        ws[ref].font = Font(name="Calibri", bold=True, size=10)
    for ref, val in (("F4", config.DEFECT_THRESHOLD), ("F5", config.WARNING_THRESHOLD)):
        ws[ref].value = val
        ws[ref].number_format = PCT_FMT
        ws[ref].fill = YEL_FILL
        ws[ref].font = Font(name="Calibri", bold=True, size=10, color="1D4ED8")

    node_crit, carrier_crit, cat_crit = "$C$4", "$C$5", "$C$6"
    w_first = 14
    w_last = w_first + len(weeks) - 1

    # KPI strip
    kp = ["Shipments (filtered)", "Defects (filtered)", "Defect Rate", "Latest Full-Week Rate",
          "Latest WoW Change", "Reconciliation vs pipeline"]
    for i, h in enumerate(kp, 1):
        _hdr(ws, 8, i, h)
    _cell(ws, 9, 1, f"=SUM(B{w_first}:B{w_last})", bold=True, fmt=INT_FMT)
    _cell(ws, 9, 2, f"=SUM(C{w_first}:C{w_last})", bold=True, fmt=INT_FMT)
    _cell(ws, 9, 3, "=IFERROR(B9/A9,0)", bold=True, fmt=PCT_FMT)
    _cell(ws, 9, 4, f'=IFERROR(LOOKUP(2,1/(E{w_first}:E{w_last}=7),D{w_first}:D{w_last}),0)',
          bold=True, fmt=PCT_FMT)
    _cell(ws, 9, 5, f'=IFERROR(LOOKUP(2,1/(F{w_first}:F{w_last}<>""),F{w_first}:F{w_last}),0)',
          bold=True, fmt='+0.00%;-0.00%;0.00%')
    _cell(ws, 9, 6, f'=IF(SUM({_rng("F", data_last)})={int(python_total_defects)},'
                    f'"MATCH ({int(python_total_defects):,} defects)","MISMATCH")', bold=True)
    ws.row_dimensions[8].height = 30

    # weekly table
    ws.cell(row=12, column=1, value="Weekly trend (filtered by the selectors above)").font = \
        Font(name="Calibri", bold=True, size=11, color=CLR["header_bg"])
    heads = ["Week start", "Shipments", "Defects", "Defect Rate", "Days", "WoW Change",
             *[DEFECT_LABEL[c] for c in PIVOT_DEFECT_ORDER], "Status"]
    for i, h in enumerate(heads, 1):
        _hdr(ws, 13, i, h)
    for k, wkd in enumerate(weeks):
        r = w_first + k
        crit = [(DW["week"], f"$A{r}"), (DW["node"], node_crit),
                (DW["carrier"], carrier_crit), (DW["category"], cat_crit)]
        _cell(ws, r, 1, pd.Timestamp(wkd).to_pydatetime(), fmt="yyyy-mm-dd", align="left")
        _cell(ws, r, 2, _sumifs(DW["ship"], data_last, crit), fmt=INT_FMT)
        _cell(ws, r, 3, _sumifs(DW["defects"], data_last, crit), fmt=INT_FMT)
        _cell(ws, r, 4, f"=IFERROR(C{r}/B{r},0)", fmt=PCT_FMT)
        _cell(ws, r, 5, f'=IFERROR(AVERAGEIFS({_rng(DW["days"], data_last)},'
                        f'{_rng(DW["week"], data_last)},$A{r}),0)')
        if k == 0:
            _cell(ws, r, 6, "")
        else:
            _cell(ws, r, 6, f'=IF(AND(E{r}=7,E{r-1}=7),D{r}-D{r-1},"")', fmt='+0.00%;-0.00%;0.00%')
        for j, dc in enumerate(PIVOT_DEFECT_ORDER):
            _cell(ws, r, 7 + j, _sumifs(DW[dc], data_last, crit), fmt=INT_FMT)
        _cell(ws, r, 13, f'=IF(E{r}<7,"partial",IF(D{r}>=$F$4,"ESCALATE",IF(D{r}>=$F$5,"WATCH","OK")))')

    # node ranking table
    n_title = w_last + 3
    ws.cell(row=n_title, column=1, value="Node ranking (filtered by Carrier and Category; Node selector ignored)").font = \
        Font(name="Calibri", bold=True, size=11, color=CLR["header_bg"])
    n_hdr = n_title + 1
    nheads = ["Node", "Shipments", "Defects", "Defect Rate", "Rank", "Status",
              *[DEFECT_LABEL[c] for c in PIVOT_DEFECT_ORDER], "Top Defect Type"]
    for i, h in enumerate(nheads, 1):
        _hdr(ws, n_hdr, i, h)
    n_first, n_last = n_hdr + 1, n_hdr + len(nodes)
    for k, node in enumerate(nodes):
        r = n_first + k
        crit = [(DW["node"], f"$A{r}"), (DW["carrier"], carrier_crit), (DW["category"], cat_crit)]
        _cell(ws, r, 1, node, bold=True, align="left")
        _cell(ws, r, 2, _sumifs(DW["ship"], data_last, crit), fmt=INT_FMT)
        _cell(ws, r, 3, _sumifs(DW["defects"], data_last, crit), fmt=INT_FMT)
        _cell(ws, r, 4, f"=IFERROR(C{r}/B{r},0)", fmt=PCT_FMT)
        _cell(ws, r, 5, f"=RANK(D{r},$D${n_first}:$D${n_last},0)")
        _cell(ws, r, 6, f'=IF(D{r}>=$F$4,"ESCALATE",IF(D{r}>=$F$5,"WATCH","OK"))')
        for j, dc in enumerate(PIVOT_DEFECT_ORDER):
            _cell(ws, r, 7 + j, _sumifs(DW[dc], data_last, crit), fmt=INT_FMT)
        _cell(ws, r, 13, f"=INDEX($G${n_hdr}:$L${n_hdr},MATCH(MAX(G{r}:L{r}),G{r}:L{r},0))", align="left")

    # conditional formatting (live, threshold-driven)
    for rng in (f"D{w_first}:D{w_last}", f"D{n_first}:D{n_last}", "C9"):
        ws.conditional_formatting.add(rng, CellIsRule(operator="greaterThanOrEqual", formula=["$F$4"], fill=RED_FILL))
        ws.conditional_formatting.add(rng, CellIsRule(operator="greaterThanOrEqual", formula=["$F$5"], fill=YEL_FILL))
    for rng in (f"M{w_first}:M{w_last}", f"F{n_first}:F{n_last}"):
        ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"ESCALATE"'], fill=RED_FILL))
        ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"WATCH"'], fill=YEL_FILL))
        ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"OK"'], fill=GRN_FILL))
    ws.conditional_formatting.add(f"M{w_first}:M{w_last}",
                                  CellIsRule(operator="equal", formula=['"partial"'], fill=SUB_FILL))

    # native chart
    ch = LineChart()
    ch.title = "Weekly defect rate (follows the selectors)"
    ch.y_axis.title = "Defect rate"
    ch.y_axis.number_format = "0.0%"
    ch.height, ch.width = 7.5, 16
    ch.add_data(Reference(ws, min_col=4, min_row=13, max_row=w_last), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=w_first, max_row=w_last))
    ch.x_axis.number_format = "mm-dd"
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ws.add_chart(ch, f"O{n_title - 6}")

    ws.freeze_panes = "A4"
    for i, w in enumerate([22, 14, 14, 16, 8, 13, 13, 13, 13, 13, 15, 13, 17], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.column_dimensions["A"].width = 26


def build_assumptions_sheet(ws, simulated: bool, extra_notes: list[str] | None = None) -> None:
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 90
    ws["A1"].value = "Assumptions, Methodology & Data Provenance"
    ws["A1"].font = Font(name="Calibri", bold=True, size=13, color=CLR["header_bg"])

    r = 3
    if simulated:
        ws.merge_cells(f"A{r}:C{r}")
        ws[f"A{r}"].value = ("DATA PROVENANCE: all data in this workbook is SIMULATED. Seeded anomalies "
                             "(PHX-7/AMZL carrier issue, SEA-2 Electronics inventory problem, 3 weather events, "
                             "Electronics/Furniture/IT Hardware supplier problems) exist to validate that the "
                             "detection logic finds them. Findings are NOT real Amazon results.")
        ws[f"A{r}"].fill = YEL_FILL
        ws[f"A{r}"].alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[r].height = 62
        r += 2

    def table(title: str, rows: list[tuple]):
        nonlocal r
        ws.cell(row=r, column=1, value=title).font = Font(name="Calibri", bold=True, size=11, color=CLR["header_bg"])
        r += 1
        for i, h in enumerate(["Parameter", "Value", "Meaning / how it is used"], 1):
            _hdr(ws, r, i, h)
        for p, v, m in rows:
            r += 1
            _cell(ws, r, 1, p, align="left")
            _cell(ws, r, 2, v)
            _cell(ws, r, 3, m, align="left")
            ws.cell(row=r, column=3).alignment = Alignment(wrap_text=True, vertical="center")
        r += 2

    table("Detection thresholds", [
        ("DEFECT_THRESHOLD", config.DEFECT_THRESHOLD, "Rate at or above this = ESCALATE"),
        ("WARNING_THRESHOLD", config.WARNING_THRESHOLD, "Rate at or above this = WATCH (calibrated to the simulated ~3.3% baseline)"),
        ("ZSCORE_FLAG", config.ZSCORE_FLAG, "Cross-node z-score flag"),
        ("ROLLING_Z_FLAG", config.ROLLING_Z_FLAG, "Node-day z-score vs previous same-weekday observations"),
        ("ROLLING_MIN_EXCESS", config.ROLLING_MIN_EXCESS, "...and at least this many points above its own baseline"),
        ("ROLLING_BASELINE_WEEKS", config.ROLLING_BASELINE_WEEKS, "Same-weekday observations in the baseline window"),
    ])
    table("Customer-impact weight and cost per defect (ASSUMPTIONS)", [
        (k, f"w={config.DEFECT_WEIGHTS[k]}  ${config.COST_PER_DEFECT[k]}",
         "Weight prioritises by customer consequence; cost converts defects to USD") for k in config.DEFECT_WEIGHTS
    ])
    table("Inventory and supplier", [
        ("SERVICE_LEVEL_Z", config.SERVICE_LEVEL_Z, "95% cycle service level; SS = z*sqrt(LT*sd_d^2 + d^2*sd_LT^2)"),
        ("EXCESS_DAYS_OF_SUPPLY", config.EXCESS_DAYS_OF_SUPPLY, "Above this = EXCESS"),
        ("STOCKOUT_MARGIN", config.STOCKOUT_MARGIN, "Lost margin per stocked-out unit = unit_cost x this"),
        ("OTIF_TARGET / OTIF_WATCH", f"{config.OTIF_TARGET} / {config.OTIF_WATCH}", "OTIF = on-time AND in-full per PO"),
        ("IN_FULL_TOLERANCE", config.IN_FULL_TOLERANCE, "Received >= this share of ordered counts as in-full"),
    ])
    table("Escalation scoring", [
        ("Score formula", "0-100", "100 x (0.5 x impact + 0.5 x deviation); impact = min(1, USD / domain cap)"),
        ("COST_CAP_USD", str(config.COST_CAP_USD), "Dollar value at which impact saturates, per domain"),
        ("P1 / P2 thresholds", f"{config.P1_SCORE} / {config.P2_SCORE}", "Score >= P1 -> P1; >= P2 -> P2; else P3"),
        ("SLA days", str(config.SLA_DAYS), "Due date = report date + SLA days"),
    ])
    if extra_notes:
        ws.cell(row=r, column=1, value="Notes").font = Font(name="Calibri", bold=True, size=11, color=CLR["header_bg"])
        for n in extra_notes:
            r += 1
            ws.merge_cells(f"A{r}:C{r}")
            ws.cell(row=r, column=1, value=n).alignment = Alignment(wrap_text=True, vertical="top")
            ws.row_dimensions[r].height = 30
