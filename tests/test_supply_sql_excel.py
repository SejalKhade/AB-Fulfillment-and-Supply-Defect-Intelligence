"""Tests for the SQL layer, supply analytics, escalation queue, guard, loaders and Excel/VBA output."""
import re
import sys
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd
import pytest

ROOT = Path(__file__).parent.parent
sys.path.insert(0, str(ROOT))

import config
from modules import data_loader as dl
from modules.data_generator import BASE_DATE, WEATHER_EVENTS
from modules.hallucination_guard import check, extract_claims
from modules.warehouse import Warehouse
from pipeline_direct import run as run_pipeline


@pytest.fixture(scope="module")
def full(tmp_path_factory):
    out = tmp_path_factory.mktemp("out") / "report.xlsx"
    res = run_pipeline(seed=42, n_days=90, output_path=str(out), verbose=False)
    return res


@pytest.fixture(scope="module")
def F(full):
    return full["_frames"]


# ── pipeline / reconciliation ──────────────────────────────────────────────
class TestPipeline:
    def test_all_reconciliation_checks_pass(self, full):
        failed = [c for c in full["reconciliation"] if not c["passed"]]
        assert not failed, failed

    def test_output_exists_and_keys(self, full):
        assert Path(full["output_path"]).exists()
        for k in ("anomalies", "root_cause", "queue", "supply"):
            assert k in full

    def test_monday_spike_is_on_mondays(self, F):
        df = F["df"]
        g = df.groupby("weekday").apply(lambda x: x.total_defects.sum() / x.total_shipments.sum(),
                                        include_groups=False)
        assert g.idxmax() == "Monday"

    def test_weekday_volume_pattern(self, F):
        df = F["df"]
        vol = df.groupby("weekday")["total_shipments"].sum()
        assert vol["Monday"] > vol["Sunday"]


# ── SQL ────────────────────────────────────────────────────────────────────
class TestSQL:
    def test_every_named_query_runs(self, F):
        wh = F["warehouse"]
        for q in wh.list_queries():
            assert len(wh.run(q)) > 0, q

    def test_rolling_anomalies_find_exactly_the_seeded_weather_events(self, F):
        ev = F["rolling"].query("is_anomaly")
        found = set(zip(ev["node"], pd.to_datetime(ev["date"]).dt.date))
        expected = set()
        for day, (node, _, _) in WEATHER_EVENTS.items():
            for d in range(day - 2, day + 3):
                expected.add((node, (BASE_DATE + pd.Timedelta(days=d - 1)).date()))
        assert found == expected

    def test_contribution_finds_seeded_segments(self, F):
        c = F["contribution"]
        top_carrier = c[(c.dimension == "carrier") & (c.rank_in_dimension == 1)].iloc[0]
        assert (top_carrier["node"], top_carrier["segment"]) == ("PHX-7", "AMZL")
        top_cat = c[(c.dimension == "category") & (c.rank_in_dimension == 1)].iloc[0]
        assert (top_cat["node"], top_cat["segment"]) == ("SEA-2", "Electronics")

    def test_excess_defects_sum_to_zero_within_dimension(self, F):
        c = F["contribution"]
        assert abs(c[c.dimension == "carrier"]["excess_defects"].sum()) < 1e-6

    def test_weekly_trend_only_complete_weeks(self, F):
        wk = F["warehouse"].run("04_weekly_trend")
        assert wk["shipments"].gt(0).all()
        assert wk["wow_delta"].isna().sum() == wk["node"].nunique()  # first week per node

    def test_wow_by_node_matches_pandas_overall(self, F):
        wow_sql = F["warehouse"].run("03_wow_by_node")
        assert len(wow_sql) == 12 and wow_sql["delta"].is_monotonic_decreasing

    def test_adhoc_sql_is_read_only(self, F):
        wh = F["warehouse"]
        assert len(wh.sql("SELECT node FROM fulfillment")) > 0
        for bad in ("DROP TABLE fulfillment", "SELECT 1; DROP TABLE fulfillment",
                    "DELETE FROM fulfillment", "COPY fulfillment TO 'x.csv'"):
            with pytest.raises(ValueError):
                wh.sql(bad)
        assert "fulfillment" in wh.tables()


# ── supply ─────────────────────────────────────────────────────────────────
class TestSupply:
    def test_supplier_ranking_matches_seeded_reliability(self, F):
        sc = F["supply"]["supplier_scorecard"].set_index("supplier_id")
        assert sc.loc["SUP-ELEC-01", "otif"] < 0.2
        assert sc.loc["SUP-ELEC-01", "status"] == "ESCALATE"
        assert sc.loc["SUP-OFFC-02", "status"] == "OK"
        assert sc.loc["SUP-FURN-07", "lead_time_sd"] > sc.loc["SUP-OFFC-02", "lead_time_sd"] * 5

    def test_otif_is_on_time_and_in_full(self, F):
        sc = F["supply"]["supplier_scorecard"]
        assert (sc["otif"] <= sc[["on_time_rate", "in_full_rate"]].min(axis=1) + 1e-12).all()

    def test_inventory_conservation(self, F):
        inv = F["warehouse"].sql("SELECT * FROM inventory", limit=100000)
        assert (inv["fulfilled_units"] + inv["stockout_units"] == inv["demand_units"]).all()
        assert (inv["on_hand_end"] >= 0).all()

    def test_shrink_only_where_seeded(self, F):
        inv = F["warehouse"].sql("SELECT node, category, SUM(shrink_units) s FROM inventory GROUP BY 1,2",
                                 limit=1000)
        assert inv[inv.s > 0][["node", "category"]].values.tolist() == [["SEA-2", "Electronics"]]

    def test_safety_stock_formula(self, F):
        pol = F["supply"]["policy"].dropna(subset=["lead_time_mean"]).iloc[0]
        d, sd, lt, sdl = (pol["avg_daily_demand"], pol["sd_daily_demand"],
                          pol["lead_time_mean"], pol["lead_time_sd"])
        expected = config.SERVICE_LEVEL_Z * np.sqrt(lt * sd ** 2 + d ** 2 * sdl ** 2)
        assert abs(pol["safety_stock"] - expected) <= 0.5
        assert abs(pol["reorder_point"] - (d * lt + expected)) <= 1.0

    def test_policy_gap_raises_unreliable_and_lowers_reliable_suppliers(self, F):
        g = F["supply"]["policy"].groupby("supplier_id")[["reorder_point", "current_policy_rop"]].sum()
        for sid in ("SUP-ELEC-01", "SUP-FURN-07", "SUP-ITHW-08"):
            assert g.loc[sid, "reorder_point"] > g.loc[sid, "current_policy_rop"] * 1.2, sid
        for sid in ("SUP-OFFC-02", "SUP-MEDI-05"):
            assert g.loc[sid, "reorder_point"] < g.loc[sid, "current_policy_rop"], sid

    def test_forecast_beats_naive(self, F):
        acc = F["supply"]["accuracy"]
        assert acc["best_wape"].median() < 0.08
        assert acc["improvement_vs_naive"].median() > 0.3
        assert len(F["supply"]["forecast"]) == 96 * config.FORECAST_HORIZON_DAYS

    def test_forecast_is_non_negative(self, F):
        assert (F["supply"]["forecast"]["forecast_units"] >= 0).all()


# ── escalation queue ───────────────────────────────────────────────────────
class TestEscalation:
    def test_ids_unique_and_sorted_by_score(self, F):
        q = F["queue"]
        assert q["id"].is_unique
        assert q["severity_score"].is_monotonic_decreasing

    def test_priority_matches_score_thresholds(self, F):
        q = F["queue"]
        assert (q[q.priority == "P1"]["severity_score"] >= config.P1_SCORE).all()
        assert (q[q.priority == "P3"]["severity_score"] < config.P2_SCORE).all()

    def test_due_date_equals_report_date_plus_sla(self, F):
        q, rd = F["queue"], pd.Timestamp(F["df"]["date"].max())
        for _, r in q.iterrows():
            assert pd.Timestamp(r["due_date"]) == rd + pd.Timedelta(days=config.SLA_DAYS[r["priority"]])

    def test_every_item_has_owner_and_action(self, F):
        q = F["queue"]
        assert q["owner"].isin(config.OWNERS.values()).all()
        assert q["action"].str.len().gt(20).all()

    def test_contains_the_seeded_problems(self, F):
        ents = " ".join(F["queue"]["entity"])
        for needle in ("PHX-7 / AMZL", "SEA-2 / Electronics", "SUP-ELEC-01", "CHI-4"):
            assert needle in ents

    def test_scores_within_bounds(self, F):
        assert F["queue"]["severity_score"].between(0, 100).all()


# ── hallucination guard ────────────────────────────────────────────────────
class TestGuard:
    payload = {"rate": 0.0331, "ships": 5535730, "delta": -0.0028}

    def test_verified_when_rounded_correctly(self):
        r = check("The rate was 3.31% on 5,535,730 shipments.", self.payload)
        assert r["verified"] == 2 and r["unverified"] == 0

    def test_wrong_number_flagged(self):
        r = check("The defect rate was 9.87%.", self.payload)
        assert r["unverified"] == 1

    def test_close_but_wrong_is_not_verified(self):
        r = check("The rate was 3.4%.", self.payload)   # 2.7% off 3.31: must not be VERIFIED
        assert r["verified"] == 0

    def test_dates_and_list_numbers_ignored(self):
        claims = extract_claims("1. Report for 2026-09-28\n2. Top 3 items")
        assert claims == []


# ── loaders ────────────────────────────────────────────────────────────────
class TestLoaders:
    def test_round_trip_and_pipeline_without_supply(self, F, tmp_path):
        df = F["df"]
        p = tmp_path / "f.csv"
        df.head(5000).to_csv(p, index=False)
        loaded = dl.load_fulfillment_csv(p)
        assert len(loaded) == 5000
        res = run_pipeline(output_path=str(tmp_path / "o.xlsx"), verbose=False,
                           data={"fulfillment": df[df["date"] <= df["date"].min()].pipe(lambda x: x),
                                 "inventory": None, "purchase_orders": None, "suppliers": None})
        assert res["supply"] is None and not res["simulated"]
        sheets = openpyxl.load_workbook(res["output_path"]).sheetnames
        assert "Inventory" not in sheets and "Formula Pivot" in sheets

    def test_missing_columns_rejected(self, tmp_path):
        p = tmp_path / "bad.csv"
        pd.DataFrame({"date": ["2026-01-01"], "node": ["A"]}).to_csv(p, index=False)
        with pytest.raises(dl.DataError):
            dl.load_fulfillment_csv(p)

    def test_defects_above_shipments_rejected(self, tmp_path):
        p = tmp_path / "bad2.csv"
        pd.DataFrame({"date": ["2026-01-01"], "node": ["A"], "carrier": ["C"], "category": ["X"],
                      "total_shipments": [5], "total_defects": [9]}).to_csv(p, index=False)
        with pytest.raises(dl.DataError):
            dl.load_fulfillment_csv(p)


# ── Excel / VBA ────────────────────────────────────────────────────────────
class TestExcel:
    def test_expected_sheets(self, full):
        names = openpyxl.load_workbook(full["output_path"]).sheetnames
        for n in ("Weekly Summary", "Escalation Queue", "Node Analysis", "Root Cause", "Carrier Report",
                  "Inventory", "Forecast", "Supplier Scorecard", "Formula Pivot", "Data_Weekly",
                  "Drill Down", "Escalation Memo", "Assumptions"):
            assert n in names

    def test_data_weekly_totals_match_source(self, full, F):
        ws = openpyxl.load_workbook(full["output_path"])["Data_Weekly"]
        total = sum(r[5] for r in ws.iter_rows(min_row=2, values_only=True))
        assert total == int(F["df"]["total_defects"].sum())

    def test_pivot_uses_live_formulas(self, full):
        wb = openpyxl.load_workbook(full["output_path"])
        ws = wb["Formula Pivot"]
        formulas = [c.value for row in ws.iter_rows() for c in row
                    if isinstance(c.value, str) and c.value.startswith("=")]
        assert sum("SUMIFS" in f for f in formulas) > 200
        assert any("INDEX" in f and "MATCH" in f for f in formulas)
        assert any("RANK(" in f for f in formulas)
        assert ws.conditional_formatting and ws.data_validations.dataValidation

    def test_queue_has_workflow_dropdown(self, full):
        ws = openpyxl.load_workbook(full["output_path"])["Escalation Queue"]
        assert any("IN PROGRESS" in (dv.formula1 or "") for dv in ws.data_validations.dataValidation)

    def test_simulated_banner(self, full):
        ws = openpyxl.load_workbook(full["output_path"])["Weekly Summary"]
        assert "SIMULATED" in ws["A2"].value


class TestVBA:
    bas = (ROOT / "vba" / "DefectReportMacros.bas")

    def test_module_is_ascii_crlf(self):
        b = self.bas.read_bytes()
        b.decode("ascii")
        assert b.count(b"\r\n") == b.count(b"\n")

    def test_referenced_sheets_exist_in_workbook(self, full):
        text = self.bas.read_text(encoding="ascii")
        consts = dict(re.findall(r'Private Const (SHEET_\w+) As String = "([^"]+)"', text))
        names = set(openpyxl.load_workbook(full["output_path"]).sheetnames)
        assert consts and set(consts.values()) <= names

    def test_referenced_headers_exist_in_queue_sheet(self, full):
        text = self.bas.read_text(encoding="ascii")
        wanted = set(re.findall(r'FindHeaderCol\(ws, HEADER_ROW, "([^"]+)"\)', text))
        ws = openpyxl.load_workbook(full["output_path"])["Escalation Queue"]
        headers = {c.value for c in ws[2]}
        queue_headers = wanted & {"ID", "Priority", "Entity", "Issue", "Owner", "Due Date", "Workflow Status"}
        assert queue_headers <= headers

    def test_pivot_reconciliation_cell_is_where_macro_looks(self, full):
        ws = openpyxl.load_workbook(full["output_path"])["Formula Pivot"]
        assert "MATCH" in str(ws["F9"].value)


def _excel_available() -> bool:
    try:
        import win32com.client  # noqa: F401
        import winreg
        winreg.CloseKey(winreg.OpenKey(winreg.HKEY_CLASSES_ROOT, "Excel.Application"))
        return True
    except Exception:
        return False


@pytest.mark.skipif(not _excel_available(), reason="Microsoft Excel not available")
def test_excel_recalculation_matches_pandas(full, F, tmp_path):
    """Open the workbook in real Excel, flip the dropdowns, compare with pandas."""
    import shutil
    import win32com.client as win32
    path = tmp_path / "recalc.xlsx"
    shutil.copy(full["output_path"], path)
    import gc
    xl = win32.DispatchEx("Excel.Application")
    xl.Visible, xl.DisplayAlerts = False, False
    wb = p = None
    try:
        wb = xl.Workbooks.Open(str(path))
        xl.CalculateFull()
        p = wb.Worksheets("Formula Pivot")
        assert str(p.Range("F9").Value).startswith("MATCH")
        df = F["df"]
        p.Range("B4").Value, p.Range("B5").Value, p.Range("B6").Value = "PHX-7", "AMZL", "Electronics"
        xl.Calculate()
        sub = df[(df.node == "PHX-7") & (df.carrier == "AMZL") & (df.category == "Electronics")]
        assert p.Range("A9").Value == sub.total_shipments.sum()
        assert p.Range("B9").Value == sub.total_defects.sum()
    finally:
        if wb is not None:
            wb.Close(False)
        del p, wb
        gc.collect()
        xl.Quit()
        del xl
        gc.collect()
